import pandas as pd
import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from copy import copy
import os
import numpy as np
import holidays
import calendar
from datetime import datetime, date

# ==========================================
# CONFIGURATION & OPTIONS
# ==========================================

ENABLE_DEBUG_INFO = True

def log_debug(message):
    if ENABLE_DEBUG_INFO:
        print(f"[DEBUG] {message}")

# Helper function to automatically handle directory vs file structures
def find_actual_file(base_path):
    if os.path.isfile(base_path):
        return base_path
    folder_path = base_path.replace('.xlsx', '').replace('.csv', '')
    if os.path.isdir(folder_path):
        log_debug(f"'{folder_path}' is a directory. Looking for data files inside...")
        valid_files = [f for f in os.listdir(folder_path) if not f.startswith('_') and not f.startswith('.')]
        if not valid_files: raise FileNotFoundError(f"No valid data files in {folder_path}")
        return os.path.join(folder_path, valid_files[0])
    raise FileNotFoundError(f"Could not find a file or directory at {base_path}")


# ==========================================
# PART 1: REGION A PAYROLL PROCESSING
# ==========================================

# File Paths for Region A
REGION_A_BASE_INPUT = "/path/to/Region_A_Payroll/Inputs"
REGION_A_BASE_OUTPUT = "/path/to/Region_A_Payroll/Outputs" 

PATH_PROVIDER_A_INPUT = f"{REGION_A_BASE_INPUT}/provider_a_data"
PATH_PROVIDER_B_INPUT = f"{REGION_A_BASE_INPUT}/provider_b_data" 
PATH_TEMPLATE_INPUT = f"{REGION_A_BASE_INPUT}/master_template"

PATH_PROVIDER_A_XLSX_FALLBACK = f"{REGION_A_BASE_INPUT}/provider_a_fallback.xlsx"
PATH_OUTPUT_TEMPLATE = f"{REGION_A_BASE_OUTPUT}/Template_Updated.xlsx"
PATH_OUTPUT_REMOTE_XLSX = f"{REGION_A_BASE_OUTPUT}/remote_export.xlsx"
PATH_OUTPUT_REMOTE_CSV = f"{REGION_A_BASE_OUTPUT}/remote_export.csv"

os.makedirs(REGION_A_BASE_OUTPUT, exist_ok=True)


# STEP 1: Process Provider A Input
actual_provider_a_file = find_actual_file(PATH_PROVIDER_A_INPUT)
log_debug(f"Reading Provider A file: {actual_provider_a_file}...")

if actual_provider_a_file.endswith('.xlsx'):
    PATH_PROVIDER_A_TO_LOAD = actual_provider_a_file
else:
    df_provider_a = None
    if actual_provider_a_file.endswith('.xls'):
        df_provider_a = pd.read_excel(actual_provider_a_file)
    elif actual_provider_a_file.endswith('.parquet'):
        df_provider_a = pd.read_parquet(actual_provider_a_file)
    else:
        for enc in ['utf-8', 'latin1', 'windows-1252', 'utf-16']:
            for sep in [';', ',', '\t', None]:
                try:
                    df_provider_a = pd.read_csv(actual_provider_a_file, sep=sep, encoding=enc, engine='python' if sep is None else 'c')
                    break 
                except: continue 
            if df_provider_a is not None: break 
    df_provider_a.to_excel(PATH_PROVIDER_A_XLSX_FALLBACK, index=False)
    PATH_PROVIDER_A_TO_LOAD = PATH_PROVIDER_A_XLSX_FALLBACK


# STEP 2: Explicitly Calculate Provider B Data
log_debug("Extracting Provider B data with strict grid coordinates...")
actual_provider_b_file = find_actual_file(PATH_PROVIDER_B_INPUT)
wb_provider_b = openpyxl.load_workbook(actual_provider_b_file, data_only=True)
ws_provider_b = wb_provider_b.active

provider_b_headers = []
for c in range(16, 28): # Columns 16 (P) to 27 (AA)
    val = ws_provider_b.cell(row=6, column=c).value
    
    # Translate specific Pay Items for the remote outputs
    if isinstance(val, str):
        if val.strip() == 'Benefit Item 1 (Exempt)':
            val = 'Local Benefit Item 1 Exento'
        elif val.strip() == 'Benefit Item 2 (Exempt)':
            val = 'Local Benefit Item 2 Exento'
            
    provider_b_headers.append(val if val else f"Unknown_Header_{c}")

extracted_data = []
for r in range(7, ws_provider_b.max_row + 1):
    id_val = ws_provider_b.cell(row=r, column=2).value # Col B
    if not id_val: continue
    
    eff_date = ws_provider_b.cell(row=r, column=14).value # Col N
    end_date = ws_provider_b.cell(row=r, column=15).value # Col O
    
    for i, c in enumerate(range(16, 28)):
        amt = ws_provider_b.cell(row=r, column=c).value
        if amt is not None:
            try:
                amt_float = float(amt)
                if amt_float > 0:
                    extracted_data.append({
                        'ID': id_val,
                        'Pay_Item': provider_b_headers[i],
                        'Amount': amt_float,
                        'Eff_Date': eff_date,
                        'End_Date': end_date
                    })
            except (ValueError, TypeError):
                pass


# STEP 3: Load Reference Tables & Match IDs
log_debug("Matching Employee IDs to Provider B data...")
actual_template_file = find_actual_file(PATH_TEMPLATE_INPUT)

try:
    df_emp_id = pd.read_excel(actual_template_file, sheet_name="Employee Data", header=0)
    col_A_emp = df_emp_id.iloc[:, 0].astype(str).fillna("").tolist() 
    col_C_emp = df_emp_id.iloc[:, 2].astype(str).fillna("").tolist() 
    col_D_emp = df_emp_id.iloc[:, 3].astype(str).fillna("").tolist() 
except Exception as e:
    col_A_emp, col_C_emp, col_D_emp = [], [], []

def get_emp_id(name):
    name_str = str(name).strip().lower()
    for c_val, d_val in zip(col_C_emp, col_D_emp):
        if not c_val or str(c_val).lower() == 'nan': continue
        if name_str in str(c_val).strip().lower() or str(c_val).strip().lower() in name_str:
            return d_val
    return ""

for d in extracted_data:
    d['Emp_ID'] = get_emp_id(d['ID'])

try:
    df_lookup_a = pd.read_excel(actual_template_file, sheet_name="Provider_A_Config", usecols="N:P", header=None)
    lookup_dict_a = dict(zip(df_lookup_a.iloc[:, 0].astype(str).str.strip(), df_lookup_a.iloc[:, 2]))
except Exception as e:
    lookup_dict_a = {}


# STEP 4: Update Provider Visual Sheets
log_debug("Updating provider sheets in the Template...")
wb_template = openpyxl.load_workbook(actual_template_file, data_only=False)
wb_provider_a = openpyxl.load_workbook(PATH_PROVIDER_A_TO_LOAD, data_only=True)
ws_template_provider_a = wb_template["Provider_A_Config"]
ws_provider_a = wb_provider_a.active
max_row_provider_a = ws_provider_a.max_row

# --- Provider A Sheet Update ---
for m in list(ws_template_provider_a.merged_cells.ranges):
    if m.min_row >= 3 and m.min_col <= 7 and m.max_col >= 2:
        ws_template_provider_a.unmerge_cells(str(m))

for row_idx in range(3, max_row_provider_a + 1):
    for col_idx in range(2, 8): 
        s_cell = ws_provider_a.cell(row=row_idx, column=col_idx)
        t_cell = ws_template_provider_a.cell(row=row_idx, column=col_idx)
        if not isinstance(s_cell, MergedCell): t_cell.value = s_cell.value
        if s_cell.has_style:
            t_cell.font = copy(s_cell.font)
            t_cell.border = copy(s_cell.border)
            t_cell.fill = copy(s_cell.fill)
            t_cell.number_format = copy(s_cell.number_format)
            t_cell.alignment = copy(s_cell.alignment)

for m_range in ws_provider_a.merged_cells.ranges:
    if m_range.min_row >= 3 and m_range.min_col >= 2 and m_range.max_col <= 7:
        try: ws_template_provider_a.merge_cells(str(m_range))
        except: pass

def match_provider_a_col_I(d_val):
    if pd.isna(d_val) or str(d_val).strip() == "": return "Not Found"
    d_str = str(d_val).lower()
    for a_val, d_out in zip(col_A_emp, col_D_emp):
        if not a_val: continue
        if str(a_val).strip().lower() in d_str: return d_out
    return "Not Found"

for r in range(5, ws_template_provider_a.max_row + 1):
    d_cell = ws_template_provider_a.cell(row=r, column=4) 
    g_cell = ws_template_provider_a.cell(row=r, column=7) 
    
    if not isinstance(d_cell, MergedCell) and d_cell.value is not None:
        ws_template_provider_a.cell(row=r, column=9).value = match_provider_a_col_I(d_cell.value)
        ws_template_provider_a.cell(row=r, column=10).value = d_cell.value
        d_str = str(d_cell.value).strip()
        g_val = pd.to_numeric(g_cell.value, errors='coerce') if g_cell.value is not None else 0
        p_val = pd.to_numeric(lookup_dict_a.get(d_str, 0), errors='coerce')
        ws_template_provider_a.cell(row=r, column=11).value = p_val * g_val
    else:
        for col in [9, 10, 11]:
            if not isinstance(ws_template_provider_a.cell(row=r, column=col), MergedCell):
                ws_template_provider_a.cell(row=r, column=col).value = None

# --- Provider B Sheet Update ---
ws_template_provider_b = wb_template["Provider_B_Config"] 
provider_b_source_rows = max(0, ws_provider_b.max_row - 5)

template_target_rows = sum(1 for r in range(7, ws_template_provider_b.max_row + 1) if not isinstance(ws_template_provider_b.cell(row=r, column=1), MergedCell) and ws_template_provider_b.cell(row=r, column=1).value is not None)

provider_b_diff = provider_b_source_rows - template_target_rows
if provider_b_diff > 0: ws_template_provider_b.insert_rows(7 + template_target_rows, amount=provider_b_diff)
elif provider_b_diff < 0: ws_template_provider_b.delete_rows(7 + provider_b_source_rows, amount=abs(provider_b_diff))

for r_offset in range(provider_b_source_rows):
    source_r = 6 + r_offset
    target_r = 7 + r_offset
    for c in range(1, 28): 
        target_cell = ws_template_provider_b.cell(row=target_r, column=c)
        if not isinstance(target_cell, MergedCell):
            target_cell.value = ws_provider_b.cell(row=source_r, column=c).value


# STEP 5: Save Updated Template
log_debug(f"Saving fully updated Template to: {PATH_OUTPUT_TEMPLATE}")
wb_template.save(PATH_OUTPUT_TEMPLATE)


# STEP 6: Safely Inject Content to remote file
log_debug("Isolating 'output remote' sheet and securely injecting text values...")

actual_remote_sheet_name = next((s for s in wb_template.sheetnames if s.strip().lower() == "output remote"), None)
if actual_remote_sheet_name is None:
    raise KeyError(f"Could not find 'output remote'. Available sheets: {wb_template.sheetnames}")

ws_remote = wb_template[actual_remote_sheet_name]
ws_remote.sheet_state = 'visible'

for r in range(5, ws_remote.max_row + 50):
    for c in range(1, 9):
        if not isinstance(ws_remote.cell(row=r, column=c), MergedCell):
            ws_remote.cell(row=r, column=c).value = None

last_r = 4

# PROVIDER B INJECTION
for idx, d in enumerate(extracted_data):
    r = 5 + idx
    ws_remote.cell(row=r, column=1).value = d['ID']
    ws_remote.cell(row=r, column=2).value = d['Pay_Item'] 
    ws_remote.cell(row=r, column=3).value = d['Amount']
    ws_remote.cell(row=r, column=4).value = "LOCAL_CURRENCY" 
    ws_remote.cell(row=r, column=5).value = "ONCE"
    ws_remote.cell(row=r, column=6).value = None
    ws_remote.cell(row=r, column=7).value = d['Eff_Date']
    ws_remote.cell(row=r, column=8).value = d['End_Date']
    last_r = r

# PROVIDER A INJECTION
log_debug("Appending Provider A data to 'output remote' sheet...")
current_r = last_r + 1
data_provider_a_appended = []

for row_idx in range(5, ws_template_provider_a.max_row + 1):
    col_i_val = ws_template_provider_a.cell(row=row_idx, column=9).value
    col_k_val = ws_template_provider_a.cell(row=row_idx, column=11).value
    
    if col_k_val is not None:
        try:
            col_k_val = round(float(col_k_val), 2)
        except (ValueError, TypeError):
            pass 
    
    if col_i_val is not None and str(col_i_val).strip() != "" and str(col_i_val).strip() != "Not Found":
        
        pay_item_val = 'Tax Exempt Code Local'
        
        ws_remote.cell(row=current_r, column=1).value = col_i_val  
        ws_remote.cell(row=current_r, column=2).value = pay_item_val             
        ws_remote.cell(row=current_r, column=3).value = col_k_val  
        ws_remote.cell(row=current_r, column=4).value = "LOCAL_CURRENCY"
        ws_remote.cell(row=current_r, column=5).value = "ONCE"
        ws_remote.cell(row=current_r, column=6).value = None
        
        val_eff_date = ws_remote.cell(row=current_r - 1, column=7).value
        val_end_date = ws_remote.cell(row=current_r - 1, column=8).value
        
        ws_remote.cell(row=current_r, column=7).value = val_eff_date
        ws_remote.cell(row=current_r, column=8).value = val_end_date
        
        data_provider_a_appended.append({
            'ID': col_i_val,
            'Emp_ID': col_i_val,
            'Pay_Item': pay_item_val,
            'Amount': col_k_val,
            'Currency': 'LOCAL_CURRENCY',
            'Occurrence': 'ONCE',
            'Eff_Date': val_eff_date,
            'End_Date': val_end_date
        })
        current_r += 1

# Isolate sheet and save
for s in wb_template.sheetnames:
    if s != actual_remote_sheet_name:
        del wb_template[s]

wb_template.save(PATH_OUTPUT_REMOTE_XLSX)
log_debug(f"Successfully saved pristine values to: {PATH_OUTPUT_REMOTE_XLSX}")


# STEP 7: Clean CSV Generation
log_debug("Building CSV directly from memory to ensure zero artifacts...")

csv_data = {
    'legacy_id_column': [d['ID'] for d in extracted_data] + [d['ID'] for d in data_provider_a_appended],
    'output_data.employee_identifier': [d['Emp_ID'] for d in extracted_data] + [d['Emp_ID'] for d in data_provider_a_appended],
    'output_data.pay_item': [d['Pay_Item'] for d in extracted_data] + [d['Pay_Item'] for d in data_provider_a_appended],
    'output_data.amount': [d['Amount'] for d in extracted_data] + [d['Amount'] for d in data_provider_a_appended],
    'output_data.currency': ["LOCAL_CURRENCY" for d in extracted_data] + [d['Currency'] for d in data_provider_a_appended],
    'output_data.occurrence': ["ONCE" for d in extracted_data] + [d['Occurrence'] for d in data_provider_a_appended],
    'output_data.effective_date': [d['Eff_Date'] for d in extracted_data] + [d['Eff_Date'] for d in data_provider_a_appended],
    'output_data.end_date': [d['End_Date'] for d in extracted_data] + [d['End_Date'] for d in data_provider_a_appended]
}

pd.DataFrame(csv_data).to_csv(PATH_OUTPUT_REMOTE_CSV, index=False)
log_debug(f"Successfully saved pristine values to: {PATH_OUTPUT_REMOTE_CSV}")


# ==========================================
# PART 2: REGION B PAYROLL PROCESSING
# ==========================================

# HELPERS: DATA FORMATTING & WORKING DAYS
def calculate_regional_working_days(target_date):
    """Calculates total weekdays (Mon-Fri) and public holidays for a given month in target region."""
    year, month = target_date.year, target_date.month
    _, last_day = calendar.monthrange(year, month)
    start_date = f"{year}-{month:02d}-01"
    end_date = f"{year}-{month:02d}-{last_day}"
    
    total_weekdays = np.busday_count(start_date, np.datetime64(end_date) + np.timedelta64(1, 'D'))
    # Replaced country-specific holiday logic with generic equivalent
    regional_holidays = holidays.CountryHoliday('FR', years=year) 
    holiday_weekdays = sum(1 for d, name in regional_holidays.items() if d.month == month and d.weekday() < 5)
    
    return int(total_weekdays), int(holiday_weekdays)

def clean_id(val):
    """Ensures IDs stay as 5-digit strings and don't get truncated."""
    if val is None: return ""
    s = str(val).strip()
    if s.endswith('.0'): s = s[:-2]
    if s.isdigit() and len(s) < 5: return s.zfill(5)
    return s

def get_excel_serial(date_val):
    """Converts a Python datetime into an Excel serial number for perfect string matching."""
    if pd.isna(date_val) or date_val is None: return ""
    if isinstance(date_val, (datetime, date)):
        return str((date_val.date() - date(1899, 12, 30)).days)
    try: return str((pd.to_datetime(date_val).date() - date(1899, 12, 30)).days)
    except: return str(date_val).strip()


# FILE PATHS
REGION_B_INPUT_BASE = "/path/to/Region_B_Payroll/Inputs"
REGION_B_OUTPUT_BASE = "/path/to/Region_B_Payroll/Outputs"
os.makedirs(REGION_B_OUTPUT_BASE, exist_ok=True)

PATH_BENEFIT_CALC_INPUT = f"{REGION_B_INPUT_BASE}/Benefit_Calculation"
PATH_CALC_PREV = f"{REGION_B_INPUT_BASE}/Calculation_Previous_Month"
PATH_HR_SYS_HOL = f"{REGION_B_INPUT_BASE}/HR_System_Holiday_Current"
PATH_HR_SYS_OTH = f"{REGION_B_INPUT_BASE}/HR_System_Absent_Other"

PATH_OUTPUT_BENEFIT = f"{REGION_B_OUTPUT_BASE}/Benefit_Calculation_updated.xlsx"
PATH_OUTPUT_REGION_B_XLSX = f"{REGION_B_OUTPUT_BASE}/remote_export.xlsx"
PATH_OUTPUT_REGION_B_CSV = f"{REGION_B_OUTPUT_BASE}/remote_export.csv"


# STEP 1: LOAD FILES (SAFE VISIBLE SHEET MAPPING)
wb_benefit = openpyxl.load_workbook(find_actual_file(PATH_BENEFIT_CALC_INPUT))
visible_benefit = [ws for ws in wb_benefit.worksheets if ws.sheet_state == 'visible']
ws1_benefit, ws2_benefit, ws3_benefit = visible_benefit[0], visible_benefit[1], visible_benefit[2] 

wb_benefit_data = openpyxl.load_workbook(find_actual_file(PATH_BENEFIT_CALC_INPUT), data_only=True)
ws3_benefit_data = [ws for ws in wb_benefit_data.worksheets if ws.sheet_state == 'visible'][2]

wb_prev_formulas = openpyxl.load_workbook(find_actual_file(PATH_CALC_PREV), data_only=False)
ws2_prev_formulas = [ws for ws in wb_prev_formulas.worksheets if ws.sheet_state == 'visible'][1]

wb_prev_data = openpyxl.load_workbook(find_actual_file(PATH_CALC_PREV), data_only=True)
ws2_prev_data = [ws for ws in wb_prev_data.worksheets if ws.sheet_state == 'visible'][1]

wb_hr_hol = openpyxl.load_workbook(find_actual_file(PATH_HR_SYS_HOL), data_only=True)
ws_hr_hol = wb_hr_hol.active

wb_hr_oth = openpyxl.load_workbook(find_actual_file(PATH_HR_SYS_OTH), data_only=True)
ws_hr_oth = wb_hr_oth.active

old_b8_dict = {}
for c in range(2, ws3_benefit_data.max_column + 1):
    old_id = ws3_benefit_data.cell(row=5, column=c).value
    old_b8 = ws3_benefit_data.cell(row=8, column=c).value
    if old_id is not None and str(old_id).strip() != "":
        old_b8_dict[clean_id(old_id)] = old_b8


# STEP 2: CALCULATION (PREV MONTH) -> BENEFIT SHEET 1
log_debug("Copying Calculation (Previous Month) Sheet 2 to Benefit Sheet 1...")
for r in range(1, ws1_benefit.max_row + 1):
    for c in range(1, 15): ws1_benefit.cell(row=r, column=c).value = None

sheet1_g_values = set()

for r in range(1, ws2_prev_formulas.max_row + 1):
    for c in range(1, 8):
        val = ws2_prev_formulas.cell(row=r, column=c).value
        
        if c == 5 and isinstance(val, str) and val.strip().lower() == 'other events - paid':
            val = 'Holiday'
            
        ws1_benefit.cell(row=r, column=c).value = val
        
        if c == 7 and r >= 2:
            val_g = ws2_prev_data.cell(row=r, column=c).value
            if val_g is not None: 
                g_str = str(val_g).strip().lower().replace('other events - paid', 'holiday')
                sheet1_g_values.add(g_str)


# STEP 3: INJECT HR DATA TO SHEET 2
log_debug("Injecting HR Holiday and Other Absent data into Sheet 2...")
for r in range(2, ws2_benefit.max_row + 1):
    for c in range(1, 15): ws2_benefit.cell(row=r, column=c).value = None

current_target_row = 2
pivot_data = {} 
all_hr_dates = [] 
safe_sheet1_name = ws1_benefit.title.replace("'", "''")

def extract_hr_to_sheet2(source_ws, start_row_target, start_src_row):
    curr_r = start_row_target
    for r_src in range(start_src_row, source_ws.max_row + 1):
        val_b = source_ws.cell(row=r_src, column=2).value 
        val_c = source_ws.cell(row=r_src, column=3).value 
        val_d = source_ws.cell(row=r_src, column=4).value 
        val_e = source_ws.cell(row=r_src, column=5).value 
        
        if not val_b: continue 
        
        if isinstance(val_e, str) and val_e.strip().lower() == 'other events - paid':
            val_e = 'Holiday'
        
        all_hr_dates.append(val_d) 
        
        ws2_benefit.cell(row=curr_r, column=1).value = val_b
        ws2_benefit.cell(row=curr_r, column=2).value = val_c
        ws2_benefit.cell(row=curr_r, column=4).value = val_d
        ws2_benefit.cell(row=curr_r, column=5).value = val_e
        
        ws2_benefit.cell(row=curr_r, column=3).value = f"=VLOOKUP(E{curr_r},Table!$B$3:$D$25,3,FALSE)"
        ws2_benefit.cell(row=curr_r, column=7).value = f"=A{curr_r}&D{curr_r}&E{curr_r}"
        ws2_benefit.cell(row=curr_r, column=6).value = f'=IF(IFERROR(VLOOKUP(G{curr_r},\'{safe_sheet1_name}\'!G:G,1,FALSE),"")=G{curr_r},"yes","no")'
        
        str_b = clean_id(val_b)
        serial_d = get_excel_serial(val_d)
        str_e = str(val_e).strip() if val_e else ""
        concat_val = f"{str_b}{serial_d}{str_e}".lower()
        
        is_yes = any(concat_val in g or concat_val == g for g in sheet1_g_values)
                
        if not is_yes:
            if str_b not in pivot_data: pivot_data[str_b] = 0
            try: pivot_data[str_b] += float(val_c) if val_c is not None else 0.0
            except (ValueError, TypeError): pass
                
        curr_r += 1
    return curr_r

current_target_row = extract_hr_to_sheet2(ws_hr_hol, current_target_row, start_src_row=2)
current_target_row = extract_hr_to_sheet2(ws_hr_oth, current_target_row, start_src_row=4)

log_debug("Copying 'no' records from Sheet 1 (Previous Month) into Sheet 2...")
for r_src in range(2, ws2_prev_data.max_row + 1):
    f_val = ws2_prev_data.cell(row=r_src, column=6).value
    
    if str(f_val).strip().lower() in ['non', 'no']:
        val_a = ws2_prev_data.cell(row=r_src, column=1).value
        val_b = ws2_prev_data.cell(row=r_src, column=2).value
        val_d = ws2_prev_data.cell(row=r_src, column=4).value
        val_e = ws2_prev_data.cell(row=r_src, column=5).value
        
        if not val_a: continue 
        
        if isinstance(val_e, str) and val_e.strip().lower() == 'other events - paid':
            val_e = 'Holiday'
        
        ws2_benefit.cell(row=current_target_row, column=1).value = val_a
        ws2_benefit.cell(row=current_target_row, column=2).value = val_b
        ws2_benefit.cell(row=current_target_row, column=4).value = val_d
        ws2_benefit.cell(row=current_target_row, column=5).value = val_e
        
        ws2_benefit.cell(row=current_target_row, column=3).value = f"=VLOOKUP(E{current_target_row},Table!$B$3:$D$25,3,FALSE)"
        ws2_benefit.cell(row=current_target_row, column=7).value = f"=A{current_target_row}&D{current_target_row}&E{current_target_row}"
        ws2_benefit.cell(row=current_target_row, column=6).value = f'=IF(IFERROR(VLOOKUP(G{current_target_row},\'{safe_sheet1_name}\'!G:G,1,FALSE),"")=G{current_target_row},"yes","no")'
        
        str_b = clean_id(val_a)
        serial_d = get_excel_serial(val_d)
        str_e = str(val_e).strip() if val_e else ""
        concat_val = f"{str_b}{serial_d}{str_e}".lower()
        
        is_yes = any(concat_val in g or concat_val == g for g in sheet1_g_values)
        if not is_yes:
            if str_b not in pivot_data: 
                pivot_data[str_b] = 0
            try: pivot_data[str_b] += float(val_b) if val_b is not None else 0.0
            except (ValueError, TypeError): pass
                
        current_target_row += 1

ws2_benefit.cell(row=13, column=9).value = "Row Labels"
ws2_benefit.cell(row=13, column=10).value = "Sum of Holiday Duration (Days)"

current_pt_row = 14
total_yes = 0
for emp_id, duration in pivot_data.items():
    ws2_benefit.cell(row=current_pt_row, column=9).value = emp_id
    ws2_benefit.cell(row=current_pt_row, column=10).value = duration
    total_yes += duration
    current_pt_row += 1

current_pt_row += 1
ws2_benefit.cell(row=current_pt_row, column=9).value = "Grand Total"
ws2_benefit.cell(row=current_pt_row, column=10).value = total_yes


# STEP 4: UPDATE THIRD SHEET (LATEST MONTH)
log_debug("Automatically detecting the LATEST month from the data...")

valid_dates = pd.to_datetime(pd.Series(all_hr_dates), dayfirst=True, errors='coerce').dropna()
if not valid_dates.empty:
    target_date = valid_dates.max()
else:
    target_date = datetime.now() 

working_days, public_holidays = calculate_regional_working_days(target_date)

last_day_of_month = calendar.monthrange(target_date.year, target_date.month)[1]
hr_date_str = f"{target_date.year}-{target_date.month:02d}-{last_day_of_month:02d}"

ws3_benefit.cell(row=1, column=1).value = "Number of working days in the month"
ws3_benefit.cell(row=1, column=2).value = working_days 
ws3_benefit.cell(row=2, column=1).value = "Number of public holidays in the month"
ws3_benefit.cell(row=2, column=2).value = public_holidays
ws3_benefit.cell(row=9, column=1).value = "Number of days absent + adjustment for previous month"
ws3_benefit.cell(row=11, column=1).value = "Expense system / half-day adjustment"
ws3_benefit.cell(row=12, column=1).value = "Benefit Items Allocated"
ws3_benefit.cell(row=16, column=1).value = "Check benefit allocation total"

safe_sheet2_name = ws2_benefit.title.replace("'", "''")

log_debug("Populating Row 5 and INDEX/MATCH lookup formulas from Employees' ID sheet...")
ws_emp_id = wb_benefit["Employee Data"]
ws_emp_id_data = wb_benefit_data["Employee Data"] 
safe_emp_sheet = ws_emp_id.title.replace("'", "''")

target_col = 2
emp_name_dict = {}

max_row_to_check = max(ws_emp_id_data.max_row, 200)

for r_src in range(2, max_row_to_check + 1):
    val_f_data = ws_emp_id_data.cell(row=r_src, column=6).value
    f_str = str(val_f_data).strip() if val_f_data is not None else ""
    
    if f_str and f_str not in ["None", "#N/A", "#VALUE!", "#REF!", "0"]:
        val_c_data = ws_emp_id_data.cell(row=r_src, column=3).value
        col_let = get_column_letter(target_col)
        
        clean_f = clean_id(f_str)
        
        ws3_benefit.cell(row=5, column=target_col).value = clean_f
        emp_name_dict[clean_f] = val_c_data
        
        ws3_benefit.cell(row=4, column=target_col).value = f"=INDEX('{safe_emp_sheet}'!$C:$C, MATCH({col_let}5, '{safe_emp_sheet}'!$F:$F, 0))"
        ws3_benefit.cell(row=6, column=target_col).value = f"=INDEX('{safe_emp_sheet}'!$D:$D, MATCH({col_let}5, '{safe_emp_sheet}'!$F:$F, 0))"
        ws3_benefit.cell(row=7, column=target_col).value = f"=INDEX('{safe_emp_sheet}'!$E:$E, MATCH({col_let}5, '{safe_emp_sheet}'!$F:$F, 0))"
        
        target_col += 1

for col in range(target_col, ws3_benefit.max_column + 2):
    for clear_row in [4, 5, 6, 7, 9, 11, 12]:
        ws3_benefit.cell(row=clear_row, column=col).value = None

for col in range(2, target_col): 
    col_let = get_column_letter(col)
    ws3_benefit.cell(row=9, column=col).value = f"=IFERROR(VLOOKUP({col_let}5,'{safe_sheet2_name}'!$I:$J,2,FALSE),0)"
    ws3_benefit.cell(row=11, column=col).value = f"=MOD({col_let}9, 1)"
    ws3_benefit.cell(row=12, column=col).value = f"=$B$1-$B$2-{col_let}8-{col_let}9-{col_let}11"

ws3_benefit.cell(row=16, column=2).value = f"=SUM(B12:{get_column_letter(target_col-1)}12)*9.25"


# STEP 5: TEMPLATE REMOTE SHEET 
log_debug("Rebuilding TEMPLATE REMOTE from scratch...")
remote_sheet_name = "OUTPUT REMOTE"

if remote_sheet_name in wb_benefit.sheetnames: del wb_benefit[remote_sheet_name]
ws_remote = wb_benefit.create_sheet(remote_sheet_name)
safe_ws3_name = ws3_benefit.title.replace("'", "''")

ws_remote.cell(row=1, column=2).value = 'output_data.employee_identifier'
ws_remote.cell(row=1, column=3).value = 'output_data.pay_item'
ws_remote.cell(row=1, column=4).value = 'output_data.amount'
ws_remote.cell(row=1, column=5).value = 'output_data.currency'
ws_remote.cell(row=1, column=6).value = 'output_data.occurrence'
ws_remote.cell(row=1, column=7).value = 'output_data.effective_date'
ws_remote.cell(row=1, column=8).value = 'output_data.end_date'

for idx in range(target_col - 2): 
    r = 2 + idx
    src_col_letter = get_column_letter(r)
    
    ws_remote.cell(row=r, column=2).value = f'=IF(\'{safe_ws3_name}\'!{src_col_letter}4="","",\'{safe_ws3_name}\'!{src_col_letter}4)'
    ws_remote.cell(row=r, column=3).value = "Benefit Allowance Item"
    ws_remote.cell(row=r, column=4).value = f'=IF(B{r}="","",\'{safe_ws3_name}\'!{src_col_letter}12)'
    ws_remote.cell(row=r, column=5).value = "Unit"
    ws_remote.cell(row=r, column=6).value = "ONCE"
    ws_remote.cell(row=r, column=7).value = f'=IF(H{r}="","",REPLACE(H{r},9,2,"01"))'
    ws_remote.cell(row=r, column=8).value = hr_date_str

log_debug(f"Saving fully updated calculation file: {PATH_OUTPUT_BENEFIT}")
wb_benefit.save(PATH_OUTPUT_BENEFIT)


# STEP 6: EXPORT (Values Only)
log_debug("Generating pure values export...")
remote_data = []

for col in range(2, target_col):
    emp_id_val = str(ws3_benefit.cell(row=5, column=col).value).strip()
    emp_name_val = emp_name_dict.get(clean_id(emp_id_val), "")
    
    if emp_name_val and str(emp_name_val).strip() not in ["", "None"]:
        emp_key = clean_id(emp_id_val)
        current_leave = pivot_data.get(emp_key, 0)
        
        b8_val_raw = old_b8_dict.get(emp_key, 0.0)
        try: b8_val = float(b8_val_raw) if b8_val_raw not in [None, '#N/A', '#REF!', '#VALUE!'] else 0.0
        except (ValueError, TypeError): b8_val = 0.0
        
        b9_val = float(current_leave)
        b11_val = b9_val % 1
        
        amt_allocated = working_days - public_holidays - b8_val - b9_val - b11_val
        col_g_calc = f"{target_date.year}-{target_date.month:02d}-01"
        
        remote_data.append({
            'B': emp_name_val,
            'C': 'Benefit Allowance Item',
            'D': amt_allocated,
            'E': 'Unit',
            'F': 'ONCE',
            'G': col_g_calc,
            'H': hr_date_str
        })

if remote_data:
    df_export = pd.DataFrame({
        'Blank_Column_A': [''] * len(remote_data),
        'output_data.employee_identifier': [row['B'] for row in remote_data],
        'output_data.pay_item': [row['C'] for row in remote_data],
        'output_data.amount': [row['D'] for row in remote_data],
        'output_data.currency': [row['E'] for row in remote_data],
        'output_data.occurrence': [row['F'] for row in remote_data],
        'output_data.effective_date': [row['G'] for row in remote_data],
        'output_data.end_date': [row['H'] for row in remote_data]
    })
    
    df_export.rename(columns={'Blank_Column_A': ''}, inplace=True)
    
    df_export.to_excel(PATH_OUTPUT_REGION_B_XLSX, index=False, header=True)
    df_export.to_csv(PATH_OUTPUT_REGION_B_CSV, index=False, header=True)
    
    log_debug(f"Successfully saved pristine values to: {PATH_OUTPUT_REGION_B_XLSX} and CSV.")
else:
    log_debug("Warning: No valid data found for export.")

log_debug("Execution complete.")
